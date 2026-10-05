"""Build the "Name It Once" PowerPoint deck for the functions topic.

Starts from the copy-paste a function removes, then walks the material in
``functions/_docs``: `def` and calling, parameters, `return` versus `print`,
type hints and docstrings, scope, the common function patterns, and imports
with the `__main__` guard -- ending with how the function exercises are graded.

The drawing primitives are the same ones used by the earlier decks, so the
whole course looks like one set.

Run:  uv run --with python-pptx python slides/functions/build_functions_deck.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).with_name("functions.pptx")

# ---------------------------------------------------------------- palette --
NAVY = RGBColor(0x1E, 0x2A, 0x38)
BLUE = RGBColor(0x37, 0x76, 0xAB)
YELLOW = RGBColor(0xFF, 0xD4, 0x3B)
GRAY = RGBColor(0x5A, 0x66, 0x72)
LIGHT = RGBColor(0xF2, 0xF5, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CODE_FG = RGBColor(0xE8, 0xEE, 0xF4)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
RED = RGBColor(0xC6, 0x28, 0x28)

SANS = "Trebuchet MS"
MONO = "Consolas"

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.7)  # left/right margin
BODY_TOP = Inches(1.75)
BODY_W = W - 2 * M

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]

_section = {"n": 0, "name": ""}


# ------------------------------------------------------------- primitives --
def _box(
    slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None
):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def _tf(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = 0
    return tf


def _para(
    tf,
    text,
    size,
    color,
    bold=False,
    font=SANS,
    first=False,
    space_before=0,
    space_after=0,
    align=PP_ALIGN.LEFT,
    italic=False,
):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    if text:
        r = p.add_run()
        r.text = text
        f = r.font
        f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, font
        f.color.rgb = color
    return p


def _rich(p, parts, size, default_color, font=SANS):
    """parts: list of (text, color|None, bold, mono)."""
    for text, color, bold, mono in parts:
        r = p.add_run()
        r.text = text
        f = r.font
        f.size = Pt(size)
        f.bold = bold
        f.name = MONO if mono else font
        f.color.rgb = color or default_color


def _markup(text):
    """`code` -> mono+blue, *bold* -> bold navy."""
    parts, buf, i = [], "", 0
    while i < len(text):
        ch = text[i]
        if ch in "`*":
            end = text.find(ch, i + 1)
            if end > i:
                if buf:
                    parts.append((buf, None, False, False))
                    buf = ""
                inner = text[i + 1 : end]
                if ch == "`":
                    parts.append((inner, BLUE, False, True))
                else:
                    parts.append((inner, NAVY, True, False))
                i = end + 1
                continue
        buf += ch
        i += 1
    if buf:
        parts.append((buf, None, False, False))
    return parts


def slide(title=None, kicker=None, footer=True):
    s = prs.slides.add_slide(BLANK)
    _box(s, 0, 0, W, H, fill=WHITE)
    if title is not None:
        _box(s, M, Inches(0.62), Inches(0.09), Inches(0.72), fill=YELLOW)
        tf = _tf(s, M + Inches(0.26), Inches(0.5), BODY_W - Inches(0.26), Inches(1.0))
        if kicker:
            _para(tf, kicker.upper(), 12, BLUE, bold=True, first=True, space_after=2)
            _para(tf, title, 30, NAVY, bold=True)
        else:
            _para(tf, title, 30, NAVY, bold=True, first=True)
    if footer and _section["name"]:
        tf = _tf(s, M, H - Inches(0.52), BODY_W, Inches(0.3))
        _para(tf, _section["name"], 10, GRAY, first=True)
    return s


def bullets(s, items, x=None, y=None, w=None, size=18, gap=9):
    """items: str, or (str, sub_level)."""
    x = M if x is None else x
    y = BODY_TOP if y is None else y
    w = BODY_W if w is None else w
    tf = _tf(s, x, y, w, H - y - Inches(0.7))
    first = True
    for it in items:
        text, lvl = (it, 0) if isinstance(it, str) else it
        if text == "":
            _para(tf, "", 6, GRAY, first=first)
            first = False
            continue
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_before = Pt(0 if p is tf.paragraphs[0] else gap)
        p.space_after = Pt(0)
        p.level = lvl
        p.alignment = PP_ALIGN.LEFT
        glyph = "▸  " if lvl == 0 else "        –  "
        r = p.add_run()
        r.text = glyph
        r.font.size = Pt(size - (2 if lvl else 0))
        r.font.name = SANS
        r.font.color.rgb = BLUE if lvl == 0 else GRAY
        r.font.bold = lvl == 0
        _rich(p, _markup(text), size - (2 if lvl else 0), NAVY if lvl == 0 else GRAY)
    return tf


def code(s, lines, x=None, y=None, w=None, size=15, caption=None, dark=True):
    x = M if x is None else x
    y = BODY_TOP if y is None else y
    w = BODY_W if w is None else w
    lh = Pt(size * 1.42).emu
    h = Emu(int(lh * len(lines)) + Inches(0.46).emu)
    _box(
        s,
        x,
        y,
        w,
        h,
        fill=NAVY if dark else LIGHT,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        radius=0.045,
    )
    tf = _tf(s, x + Inches(0.22), y + Inches(0.2), w - Inches(0.4), h - Inches(0.3))
    for i, ln in enumerate(lines):
        col = CODE_FG if dark else NAVY
        if ln.strip().startswith("#"):
            col = RGBColor(0x8E, 0xA6, 0xBD) if dark else GRAY
        _para(
            tf, ln if ln else " ", size, col, font=MONO, first=(i == 0), space_after=0
        )
    if caption:
        ctf = _tf(s, x, y + h + Inches(0.06), w, Inches(0.3))
        _para(ctf, caption, 13, GRAY, first=True, italic=True)
    return y + h + (Inches(0.42) if caption else Inches(0.24))


def out(s, lines, x=None, y=None, w=None, size=15, label="Output"):
    x = M if x is None else x
    w = BODY_W if w is None else w
    lh = Pt(size * 1.42).emu
    h = Emu(int(lh * len(lines)) + Inches(0.62).emu)
    _box(
        s,
        x,
        y,
        w,
        h,
        fill=LIGHT,
        line=RGBColor(0xD5, 0xDE, 0xE6),
        shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        radius=0.045,
    )
    tf = _tf(s, x + Inches(0.22), y + Inches(0.16), w - Inches(0.4), h - Inches(0.26))
    _para(tf, label.upper(), 10, GRAY, bold=True, first=True, space_after=3)
    for ln in lines:
        _para(tf, ln if ln else " ", size, NAVY, font=MONO)
    return y + h + Inches(0.24)


def table(
    s,
    headers,
    rows,
    x=None,
    y=None,
    w=None,
    size=14,
    col_w=None,
    mono_cols=(),
    height=None,
):
    x = M if x is None else x
    y = BODY_TOP if y is None else y
    w = BODY_W if w is None else w
    h = height or Inches(0.42 + 0.36 * len(rows))
    shp = s.shapes.add_table(len(rows) + 1, len(headers), x, y, w, h)
    tbl = shp.table
    tbl.first_row = True
    if col_w:
        total = sum(col_w)
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Emu(int(w * cw / total))
    for j, htxt in enumerate(headers):
        c = tbl.cell(0, j)
        c.text = ""
        c.fill.solid()
        c.fill.fore_color.rgb = NAVY
        c.margin_left = c.margin_right = Inches(0.1)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        _para(c.text_frame, htxt, size, WHITE, bold=True, first=True)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.text = ""
            c.fill.solid()
            c.fill.fore_color.rgb = WHITE if i % 2 else LIGHT
            c.margin_left = c.margin_right = Inches(0.1)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]
            mono = j in mono_cols
            _rich(
                p, _markup(val) if not mono else [(val, BLUE, False, True)], size, NAVY
            )
    return y + h + Inches(0.2)


def note(s, text, y, kind="tip"):
    colors = {
        "tip": (BLUE, "TIP"),
        "warn": (RED, "WATCH OUT"),
        "ok": (GREEN, "REMEMBER"),
    }
    col, label = colors[kind]
    h = Inches(0.78)
    _box(s, M, y, BODY_W, h, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    _box(s, M, y, Inches(0.07), h, fill=col)
    tf = _tf(
        s,
        M + Inches(0.26),
        y + Inches(0.1),
        BODY_W - Inches(0.5),
        h - Inches(0.2),
        anchor=MSO_ANCHOR.MIDDLE,
    )
    _para(tf, label, 10, col, bold=True, first=True, space_after=2)
    p = tf.add_paragraph()
    _rich(p, _markup(text), 15, NAVY)
    return y + h + Inches(0.2)


def section(name, subtitle, points):
    _section["n"] += 1
    _section["name"] = name
    s = prs.slides.add_slide(BLANK)
    _box(s, 0, 0, W, H, fill=NAVY)
    _box(s, 0, H - Inches(0.22), W, Inches(0.22), fill=YELLOW)
    tf = _tf(s, M, Inches(2.1), Inches(8.4), Inches(3.0))
    _para(tf, f"PART {_section['n']}", 16, YELLOW, bold=True, first=True, space_after=8)
    _para(tf, name, 44, WHITE, bold=True, space_after=10)
    _para(tf, subtitle, 19, RGBColor(0xA9, 0xBD, 0xD1))
    tf2 = _tf(s, Inches(9.0), Inches(2.3), Inches(3.6), Inches(3.0))
    _para(tf2, "IN THIS PART", 11, YELLOW, bold=True, first=True, space_after=10)
    for pt in points:
        _para(tf2, "•  " + pt, 15, RGBColor(0xD6, 0xE2, 0xEC), space_after=7)
    return s


# ================================================================== TITLE ==
def title_slide():
    s = prs.slides.add_slide(BLANK)
    _box(s, 0, 0, W, H, fill=NAVY)
    _box(s, 0, H - Inches(0.28), W, Inches(0.28), fill=YELLOW)
    _box(s, M, Inches(1.5), Inches(0.12), Inches(1.5), fill=YELLOW)
    tf = _tf(s, M + Inches(0.4), Inches(1.4), Inches(9.5), Inches(2.0))
    _para(tf, "PYTHON · TOPIC 4", 15, YELLOW, bold=True, first=True, space_after=12)
    _para(tf, "Name It Once", 52, WHITE, bold=True, space_after=10)
    _para(
        tf,
        "Functions: def, parameters, return, and modules",
        22,
        RGBColor(0xA9, 0xBD, 0xD1),
    )
    tf2 = _tf(s, M + Inches(0.4), Inches(4.3), Inches(10.5), Inches(2.0))
    _para(
        tf2,
        "A loop repeats work in one place.",
        18,
        RGBColor(0x8F, 0xB6, 0xD9),
        first=True,
        space_after=8,
    )
    _para(
        tf2,
        "A function lets you reuse it anywhere — by name.",
        18,
        RGBColor(0x8F, 0xB6, 0xD9),
    )
    tf3 = _tf(s, M + Inches(0.4), Inches(6.2), Inches(10.0), Inches(0.5))
    _para(
        tf3,
        "python-workbook2026  ·  functions",
        13,
        RGBColor(0x6E, 0x8A, 0xA3),
        first=True,
    )


def roadmap_slide():
    s = slide("What we cover today", kicker="Roadmap", footer=False)
    table(
        s,
        ["Part", "Topic", "You will be able to"],
        [
            ["1", "*Why functions*", "define a function and call it"],
            ["2", "*Parameters*", "pass values in, with defaults"],
            ["3", "*`return`*", "hand a value back — and know it is not `print`"],
            ["4", "*Type hints & docstrings*", "say what goes in and what comes out"],
            ["5", "*Scope*", "tell local variables from global ones"],
            ["6", "*Patterns*", "recognise the functions you will write all year"],
            ["7", "*Modules & imports*", "use `math`, and guard code with `__main__`"],
        ],
        y=BODY_TOP,
        col_w=(0.8, 3.6, 6.8),
        size=16,
    )
    note(
        s,
        "From today an exercise is no longer a script that prints — it is a "
        "*function* the tests call with their own arguments.",
        Inches(5.6),
        kind="ok",
    )


# =================================================================== PART 1 ==
def part1():
    section(
        "Why functions",
        "Write it once, use it many times",
        ["Copy-paste again", "`def` and the body", "Calling a function"],
    )

    s = slide("The same three lines, three times", kicker="Without a function")
    y = code(
        s,
        [
            "a = 3",
            "print(a * a)",
            "b = 7",
            "print(b * b)",
            "c = 12",
            "print(c * c)",
        ],
        w=Inches(5.6),
    )
    code(
        s,
        [
            "def square(x):",
            "    return x * x",
            "",
            "print(square(3))",
            "print(square(7))",
            "print(square(12))",
        ],
        x=M + Inches(6.1),
        y=BODY_TOP,
        w=Inches(5.6),
    )
    bullets(
        s,
        [
            "Fix a bug in `square` once — every call is fixed",
            "The name says *what* happens; the body says *how*",
        ],
        y=y + Inches(0.2),
        size=18,
    )
    note(
        s,
        "Functions you already use: `print`, `input`, `len`, `int`, `round`. "
        "Today you write your own.",
        Inches(5.9),
        kind="tip",
    )

    s = slide("Anatomy of a function", kicker="def")
    y = code(
        s,
        [
            "def greet(name):            # def  name  (parameters)  :",
            '    message = f"Hello, {name}!"',
            "    return message          # the indented block is the body",
        ],
    )
    table(
        s,
        ["Part", "Meaning"],
        [
            ["`def`", "the keyword that starts a definition"],
            ["`greet`", "the name — same rules as a variable name"],
            ["`(name)`", "the parameters — the values the function needs"],
            ["`:` + indent", "the body, exactly like `if` and `for`"],
        ],
        y=y + Inches(0.1),
        col_w=(2.4, 8.8),
        size=16,
    )
    note(
        s,
        "Defining a function *runs nothing*. The body only runs when the "
        "function is called.",
        Inches(5.9),
        kind="warn",
    )

    s = slide("Calling it", kicker="Call")
    y = code(
        s,
        [
            "def greet():",
            '    print("Hello, World!")',
            "",
            "greet()",
            "greet()",
        ],
        w=Inches(6.4),
    )
    out(
        s,
        ["Hello, World!", "Hello, World!"],
        x=M + Inches(6.8),
        y=BODY_TOP,
        w=Inches(4.8),
    )
    bullets(
        s,
        [
            "A call is the name followed by parentheses: `greet()`",
            "`greet` without `()` does not call it — it is just the function itself",
            "Define first, call after: Python reads the file top to bottom",
        ],
        y=y + Inches(0.2),
        size=18,
    )


# =================================================================== PART 2 ==
def part2():
    section(
        "Parameters",
        "Values going into a function",
        ["Parameter vs argument", "Several parameters", "Default values"],
    )

    s = slide("Parameter and argument", kicker="Inputs")
    y = code(
        s,
        [
            "def greet(name):            # name  -> parameter",
            '    print(f"Hello, {name}!")',
            "",
            'greet("Alice")              # "Alice" -> argument',
            'greet("Bob")',
        ],
        w=Inches(7.0),
    )
    out(
        s,
        ["Hello, Alice!", "Hello, Bob!"],
        x=M + Inches(7.4),
        y=BODY_TOP,
        w=Inches(4.2),
    )
    table(
        s,
        ["Word", "Where it lives", "Example"],
        [
            ["*parameter*", "in the definition — a placeholder", "`name`"],
            ["*argument*", "in the call — the real value", '`"Alice"`'],
        ],
        y=y + Inches(0.2),
        col_w=(2.4, 5.6, 3.2),
        size=16,
    )

    s = slide("Several parameters", kicker="Order matters")
    y = code(
        s,
        [
            "def rectangle_area(length, width):",
            "    return length * width",
            "",
            "print(rectangle_area(5, 3))",
            "print(rectangle_area(width=3, length=5))   # by name",
        ],
    )
    y = out(s, ["15", "15"], y=y)
    note(
        s,
        "Arguments are matched to parameters *by position*: the first value "
        "goes to the first parameter. Naming them removes the doubt.",
        y,
        kind="tip",
    )

    s = slide("Default values", kicker="Optional arguments")
    y = code(
        s,
        [
            'def greet(name, greeting="Hello"):',
            '    print(f"{greeting}, {name}!")',
            "",
            'greet("Alice")',
            'greet("Bob", "Hi")',
        ],
        w=Inches(7.0),
    )
    out(
        s,
        ["Hello, Alice!", "Hi, Bob!"],
        x=M + Inches(7.4),
        y=BODY_TOP,
        w=Inches(4.2),
    )
    note(
        s,
        "Parameters with a default come *after* the ones without. "
        '`def greet(greeting="Hello", name):` is a SyntaxError.',
        y + Inches(0.2),
        kind="warn",
    )


# =================================================================== PART 3 ==
def part3():
    section(
        "return",
        "Values coming out of a function",
        ["`return` vs `print`", "`None`", "Returning early", "Two values at once"],
    )

    s = slide("return hands the value back", kicker="return")
    y = code(
        s,
        [
            "def add(a, b):",
            "    return a + b",
            "",
            "result = add(3, 5)",
            "total = add(10, 20) + add(5, 3)",
            "print(result, total)",
        ],
        w=Inches(6.4),
    )
    out(s, ["8 38"], x=M + Inches(6.8), y=BODY_TOP, w=Inches(4.8))
    note(
        s,
        "A call with `return` is an *expression*: store it, add it, pass it to "
        "another function — like any value.",
        y + Inches(0.1),
        kind="ok",
    )

    s = slide("print is not return", kicker="The big one")
    y = code(
        s,
        [
            "def add_and_print(a, b):",
            "    print(a + b)",
            "",
            "r = add_and_print(3, 5)",
            "print(r)",
        ],
        w=Inches(5.6),
    )
    code(
        s,
        [
            "def add_and_return(a, b):",
            "    return a + b",
            "",
            "r = add_and_return(3, 5)",
            "print(r)",
        ],
        x=M + Inches(6.1),
        y=BODY_TOP,
        w=Inches(5.6),
    )
    out(s, ["8", "None"], y=y, w=Inches(5.6))
    out(s, ["8"], x=M + Inches(6.1), y=y, w=Inches(5.6))
    note(
        s,
        "A function with no `return` returns `None`. The tests check what "
        "you *return* — printing the right answer scores nothing.",
        Inches(5.75),
        kind="warn",
    )

    s = slide("return ends the function", kicker="Early return")
    y = code(
        s,
        [
            "def is_even(n):",
            "    if n % 2 == 0:",
            "        return True        # leaves right here",
            "    return False",
            "",
            "# the same, shorter: the comparison already is a bool",
            "def is_even(n):",
            "    return n % 2 == 0",
        ],
    )
    note(
        s,
        "Nothing after a `return` in the same branch ever runs — use it to "
        "leave as soon as you know the answer.",
        y,
        kind="tip",
    )

    s = slide("Returning two values", kicker="Tuples")
    y = code(
        s,
        [
            "import math",
            "",
            "def circle(radius):",
            "    area = math.pi * radius ** 2",
            "    length = 2 * math.pi * radius",
            "    return area, length",
            "",
            "a, c = circle(5)",
            'print(f"{a:.2f} {c:.2f}")',
        ],
        w=Inches(6.8),
    )
    out(s, ["78.54 31.42"], x=M + Inches(7.2), y=BODY_TOP, w=Inches(4.4))
    note(
        s,
        "`return area, length` returns a *tuple*; `a, c = ...` unpacks it "
        "into two variables.",
        y + Inches(0.1),
        kind="ok",
    )


# =================================================================== PART 4 ==
def part4():
    section(
        "Type hints & docstrings",
        "Saying what goes in and what comes out",
        ["`x: int` and `-> float`", "Docstrings", "`help()`"],
    )

    s = slide("Type hints", kicker="Annotations")
    y = code(
        s,
        [
            "def bmi(weight: float, height: float) -> float:",
            "    return weight / height ** 2",
            "",
            "def greet(name: str) -> None:",
            '    print(f"Hello, {name}!")',
        ],
    )
    y = bullets(
        s,
        [
            "`name: type` after a parameter, `-> type` after the parentheses",
            "`-> None` means: this function returns nothing",
            "VS Code reads them and warns you before you run the code",
        ],
        y=y + Inches(0.1),
        size=18,
    )
    note(
        s,
        "Hints are *documentation*. Python does not check them at run time — "
        "`bmi('70', 1.75)` still runs, and still crashes.",
        Inches(5.8),
        kind="warn",
    )

    s = slide("Docstrings", kicker="Documentation")
    y = code(
        s,
        [
            "def square(x: int) -> int:",
            '    """Return the square of x."""',
            "    return x * x",
            "",
            "help(square)",
        ],
        w=Inches(6.4),
    )
    out(
        s,
        ["square(x: int) -> int", "    Return the square of x."],
        x=M + Inches(6.8),
        y=BODY_TOP,
        w=Inches(4.8),
    )
    note(
        s,
        "A docstring is the first line of the body, in triple quotes. It says "
        "*what* the function does — not how.",
        y + Inches(0.1),
        kind="ok",
    )


# =================================================================== PART 5 ==
def part5():
    section(
        "Scope",
        "Where a variable can be seen",
        ["Local variables", "Global variables", "Shadowing"],
    )

    s = slide("What happens in a function stays there", kicker="Local")
    y = code(
        s,
        [
            "def calculate():",
            "    x = 10              # local: lives only during the call",
            "    return x * 2",
            "",
            "print(calculate())",
            "print(x)                # NameError: name 'x' is not defined",
        ],
    )
    note(
        s,
        "Parameters are local too. Inside, a function sees its own variables; "
        "outside, nobody sees them.",
        y,
        kind="ok",
    )

    s = slide("Same name, two variables", kicker="Shadowing")
    y = code(
        s,
        [
            "x = 100                 # global",
            "",
            "def calculate():",
            "    x = 10              # a new, local x",
            "    return x * 2",
            "",
            "print(calculate())",
            "print(x)",
        ],
        w=Inches(6.4),
    )
    out(s, ["20", "100"], x=M + Inches(6.8), y=BODY_TOP, w=Inches(4.8))
    note(
        s,
        "`global` lets a function change an outer variable — avoid it. Pass "
        "values in as parameters and `return` the result.",
        y + Inches(0.1),
        kind="tip",
    )


# =================================================================== PART 6 ==
def part6():
    section(
        "Patterns",
        "Five kinds of function you will write all year",
        ["Formula", "Yes / no", "Conversion", "Accumulation", "Validation"],
    )

    s = slide("Formula and yes / no functions", kicker="Patterns")
    y = code(
        s,
        [
            "import math",
            "",
            "def get_hypotenuse(a: float, b: float) -> float:",
            "    return math.sqrt(a ** 2 + b ** 2)",
            "",
            "def is_leap_year(year: int) -> bool:",
            "    return year % 4 == 0 and year % 100 != 0 or year % 400 == 0",
        ],
    )
    note(
        s,
        "Functions that answer yes / no return a `bool` and start with "
        "`is_` or `has_`: `is_even`, `is_prime`, `is_triangle`.",
        y,
        kind="ok",
    )

    s = slide("A loop inside a function", kicker="Accumulation")
    y = code(
        s,
        [
            "def digit_sum(number: int) -> int:",
            "    total = 0",
            "    number = abs(number)",
            "    while number > 0:",
            "        total += number % 10",
            "        number //= 10",
            "    return total             # after the loop",
            "",
            "print(digit_sum(12345))",
        ],
        w=Inches(7.4),
    )
    out(s, ["15"], x=M + Inches(7.8), y=BODY_TOP, w=Inches(3.8))
    note(
        s,
        "A `return` inside the loop leaves on the first turn. Indent it at the "
        "level of the loop, not the body.",
        y + Inches(0.1),
        kind="warn",
    )

    s = slide("Patterns at a glance", kicker="Summary")
    table(
        s,
        ["Pattern", "Returns", "Exercises"],
        [
            ["formula", "a number", "`circle_area`, `bmi`, `triangle_area`"],
            ["yes / no", "`True` / `False`", "`is_even`, `is_prime`, `is_leap_year`"],
            [
                "conversion",
                "the same value, new form",
                "`fahrenheit_to_celsius`, `hex_decimal`",
            ],
            ["accumulation", "a total or a string", "`digit_sum`, `caesar_cipher`"],
            ["validation", "`True` / `False`", "`is_triangle`, `password_strength`"],
        ],
        col_w=(2.4, 3.4, 5.4),
        size=16,
    )
    note(
        s,
        "Before writing the body, write the first line: name, parameters, "
        "return type. Half the problem is solved there.",
        Inches(4.9),
        kind="tip",
    )


# =================================================================== PART 7 ==
def part7():
    section(
        "Modules & imports",
        "Using code from other files",
        ["`import math`", "`from ... import ...`", '`if __name__ == "__main__"`'],
    )

    s = slide("Three ways to import", kicker="import")
    table(
        s,
        ["Write", "Then use", "When"],
        [
            ["`import math`", "`math.sqrt(16)`", "default — clear where it comes from"],
            ["`from math import sqrt, pi`", "`sqrt(16)`", "a few names, used often"],
            ["`import random as rnd`", "`rnd.randint(1, 6)`", "long module names"],
            ["`from math import *`", "`sqrt(16)`", "avoid — you lose track of names"],
        ],
        col_w=(3.8, 3.0, 4.4),
        size=16,
    )
    note(
        s,
        "Imports go at the top of the file. Never name your own file "
        "`math.py` or `random.py` — it hides the real module.",
        Inches(4.4),
        kind="warn",
    )

    s = slide('if __name__ == "__main__":', kicker="The guard")
    y = code(
        s,
        [
            "def is_even(n: int) -> bool:",
            "    return n % 2 == 0",
            "",
            'if __name__ == "__main__":',
            "    print(is_even(4))       # only when you run the file",
            "    print(is_even(7))",
        ],
    )
    y = bullets(
        s,
        [
            '`uv run is_even.py` → `__name__` is `"__main__"`, the block runs',
            "the tests do `from is_even import is_even` → the block is skipped",
        ],
        y=y + Inches(0.05),
        size=18,
    )
    note(
        s,
        "Put your own checks — and any `input()` — under the guard. Code at "
        "the top level runs on import and can break the tests.",
        Inches(5.8),
        kind="ok",
    )


# ========================================================= HOW IT'S GRADED ==
def grading():
    _section["name"] = "How the exercises work"
    s = slide("How a function exercise is graded", kicker="Workflow")
    y = code(
        s,
        [
            "# functions/is_even/is_even.py   <- you create this file",
            "def is_even(n: int) -> bool:",
            "    return n % 2 == 0",
        ],
    )
    y = bullets(
        s,
        [
            "Read `is_even.md`: it gives the function's *name* and *parameters*",
            "The test imports your function and calls it with its own arguments",
            "Name the file and the function exactly as in the statement",
            "No `input()` needed — values arrive as arguments",
        ],
        y=y + Inches(0.05),
        size=18,
    )
    out(
        s,
        ["uv run pytest functions/is_even/"],
        y=Inches(5.75),
        label="Check it",
    )


# ================================================================= CLOSING ==
def closing():
    s = prs.slides.add_slide(BLANK)
    _box(s, 0, 0, W, H, fill=NAVY)
    _box(s, 0, H - Inches(0.28), W, Inches(0.28), fill=YELLOW)
    tf = _tf(s, Inches(1.3), Inches(1.5), Inches(10.7), Inches(3.6))
    _para(tf, "Your turn", 44, WHITE, bold=True, first=True, space_after=18)
    _para(
        tf,
        "Start with is_even and circle_area, then work down the table of contents.",
        22,
        RGBColor(0xC3, 0xD3, 0xE1),
        space_after=26,
    )
    _para(
        tf,
        "Before you write a function, answer three questions: what goes in, "
        "what comes out, and what is its name.",
        17,
        RGBColor(0x8F, 0xB6, 0xD9),
    )
    tf2 = _tf(s, Inches(1.3), Inches(5.5), Inches(10.7), Inches(1.2))
    _para(
        tf2,
        "uv run pytest functions/",
        18,
        YELLOW,
        bold=True,
        first=True,
        space_after=10,
    )
    _para(tf2, "Good luck  ·  see you soon", 20, YELLOW, bold=True)


# ==================================================================== main ==
title_slide()
roadmap_slide()
part1()
part2()
part3()
part4()
part5()
part6()
part7()
grading()
closing()

prs.save(OUT)
print(f"{len(prs.slides)} slides -> {OUT}")
